# Feature 078 (SAM-KIT-UIDESC): preenche o template oficial "Application UI Description" da Samsung com a descrição do
# CCPlayTV (variante base), em inglês. Entrada: o template e as capturas de `ui-description-shots.mjs`, ambos em
# `release/samsung/kit/ui-description/`. Saída: `CCPlayTV_UI_Description_<versão>.pptx` na mesma pasta.
#
# A credencial da lista de teste NUNCA entra aqui: a tabela do caso de uso "Add a Playlist" sai com marcadores
# (<TEST_HOST>…); quem envia preenche numa cópia fora do git, na hora da submissão (release/samsung/README.md).
#
# Uso (python-pptx e Pillow num venv; não são dependências do projeto):
#   python tv-web/scripts/release/ui-description-pptx.py [--version 0.2.1]
import argparse
import copy
import tempfile
from pathlib import Path

from PIL import Image
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Emu, Inches, Pt

ROOT = Path(__file__).resolve().parents[3]
KIT = ROOT / 'release' / 'samsung' / 'kit' / 'ui-description'
TEMPLATE = KIT / 'app_description_template_eng_v142_edited3.pptx'
BUILD_DIR = Path(tempfile.mkdtemp(prefix='ccplay-uidesc-'))  # capturas reduzidas: fora do kit

TEXT = RGBColor(0x26, 0x26, 0x26)
NAVY = RGBColor(0x34, 0x49, 0x7A)
BOX_FILL = RGBColor(0xDC, 0xE6, 0xF1)
BOX_LINE = RGBColor(0x5B, 0x7B, 0xAE)
PLACEHOLDER = RGBColor(0xC0, 0x39, 0x2B)

# ---------------------------------------------------------------- utilitários

def slide_id_list(prs):
    return prs.slides._sldIdLst


def delete_slide(prs, slide):
    lst = slide_id_list(prs)
    for sld_id in list(lst):
        if prs.slides.part.related_part(sld_id.rId) is slide.part:
            prs.part.drop_rel(sld_id.rId)
            lst.remove(sld_id)
            return
    raise ValueError('slide não encontrado')


def reorder(prs, slides_in_order):
    lst = slide_id_list(prs)
    by_part = {prs.slides.part.related_part(s.rId): s for s in lst}
    wanted = [by_part[slide.part] for slide in slides_in_order]
    for el in list(lst):
        lst.remove(el)
    for el in wanted:
        lst.append(el)


def shape_named(slide, name, placeholder_ok=True):
    return next(s for s in slide.shapes if s.name == name and (placeholder_ok or not s.is_placeholder))


def latin_wrap(paragraph):
    """O template é coreano: sem isto, o PowerPoint quebra palavras latinas no meio ("chann-els")."""
    ppr = paragraph._p.get_or_add_pPr()
    ppr.set('latinLnBrk', '0')
    ppr.set('eaLnBrk', '1')


def placeholder(slide, idx):
    return next(s for s in slide.placeholders if s.placeholder_format.idx == idx)


def remove_shape(shape):
    shape._element.getparent().remove(shape._element)


def clone_into(shape, slide):
    el = copy.deepcopy(shape._element)
    slide.shapes._spTree.append(el)
    return slide.shapes[-1]


def write(tf, lines, size=None, color=TEXT, bold_first=False):
    """Reescreve um text_frame. `lines`: str ou (str, nível). Mantém o pPr do 1º parágrafo (marcador do template)."""
    paragraphs = tf.paragraphs
    first_ppr = paragraphs[0]._p.find(qn('a:pPr'))
    for p in paragraphs[1:]:
        p._p.getparent().remove(p._p)
    p0 = tf.paragraphs[0]
    for r in list(p0._p):
        if r.tag != qn('a:pPr'):
            p0._p.remove(r)
    for i, line in enumerate(lines):
        text, level = (line, 0) if isinstance(line, str) else line
        p = p0 if i == 0 else tf.add_paragraph()
        if i > 0 and first_ppr is not None:
            p._p.insert(0, copy.deepcopy(first_ppr))
        p.level = level
        latin_wrap(p)
        run = p.add_run()
        run.text = text
        run.font.color.rgb = color
        if size:
            run.font.size = Pt(size)
        if bold_first and i == 0:
            run.font.bold = True


def set_cell(cell, text, size=9, color=TEXT, bold=False):
    tf = cell.text_frame
    for p in tf.paragraphs[1:]:
        p._p.getparent().remove(p._p)
    p = tf.paragraphs[0]
    for r in list(p._p):
        if r.tag in (qn('a:r'), qn('a:br'), qn('a:fld')):
            p._p.remove(r)
    lines = text.split('\n')
    for i, line in enumerate(lines):
        para = p if i == 0 else tf.add_paragraph()
        latin_wrap(para)
        run = para.add_run()
        run.text = line
        run.font.name = 'Arial'
        run.font.size = Pt(size)
        run.font.color.rgb = color
        run.font.bold = bold


def resize_rows(table, n_body):
    """Deixa a tabela com 1 cabeçalho + `n_body` linhas, clonando a 2ª linha (ou cortando as sobras)."""
    tbl = table._tbl
    rows = tbl.findall(qn('a:tr'))
    body = rows[1:]
    while len(body) > n_body:
        tbl.remove(body.pop())
    while len(body) < n_body:
        new = copy.deepcopy(body[len(body) % 2] if len(body) > 1 else body[0])
        tbl.append(new)
        body.append(new)


def fill_table(shape, header, rows, size=9, header_size=10, row_height=None, col_widths=None):
    table = shape.table
    resize_rows(table, len(rows))
    if col_widths:
        for col, width in zip(table.columns, col_widths):
            col.width = Inches(width)
    for c, text in enumerate(header):
        set_cell(table.cell(0, c), text, size=header_size, color=RGBColor(0xFF, 0xFF, 0xFF), bold=True)
    for r, row in enumerate(rows, start=1):
        for c, text in enumerate(row):
            color = PLACEHOLDER if text.startswith('<') else TEXT
            set_cell(table.cell(r, c), text, size=size, color=color)
        if row_height:
            table.rows[r].height = Inches(row_height)
    shape.width = sum(col.width for col in table.columns)


def picture_fit(slide, path, left, top, width):
    """Insere a captura reduzida a 1280 px (o pptx fica leve) com a largura pedida."""
    out = BUILD_DIR / (Path(path).stem + '.jpg')
    with Image.open(path) as img:
        img.convert('RGB').resize((1280, 720), Image.LANCZOS).save(out, quality=86)
    return slide.shapes.add_picture(str(out), Inches(left), Inches(top), width=Inches(width))


def arrow(slide, x1, y1, x2, y2):
    conn = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(x1), Inches(y1), Inches(x2), Inches(y2))
    conn.line.color.rgb = BOX_LINE
    conn.line.width = Pt(1.25)
    ln = conn.line._get_or_add_ln()
    tail = ln.makeelement(qn('a:tailEnd'), {'type': 'triangle', 'w': 'med', 'len': 'med'})
    ln.append(tail)
    return conn


def box(slide, x, y, w, h, title, detail=None, fill=BOX_FILL):
    shp = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    shp.adjustments[0] = 0.18
    shp.fill.solid()
    shp.fill.fore_color.rgb = fill
    shp.line.color.rgb = BOX_LINE
    shp.line.width = Pt(1)
    shp.shadow.inherit = False
    tf = shp.text_frame
    tf.margin_left = tf.margin_right = Inches(0.04)
    tf.margin_top = tf.margin_bottom = Inches(0.02)
    tf.word_wrap = True
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    latin_wrap(p)
    run = p.add_run()
    run.text = title
    run.font.size = Pt(9)
    run.font.bold = True
    run.font.color.rgb = NAVY
    if detail:
        q = tf.add_paragraph()
        q.alignment = PP_ALIGN.CENTER
        latin_wrap(q)
        r2 = q.add_run()
        r2.text = detail
        r2.font.size = Pt(7.5)
        r2.font.color.rgb = TEXT
    return (x, y, w, h)


def right(b):
    return b[0] + b[2], b[1] + b[3] / 2


def left(b):
    return b[0], b[1] + b[3] / 2


def top(b):
    return b[0] + b[2] / 2, b[1]


def bottom(b):
    return b[0] + b[2] / 2, b[1] + b[3]


def link(slide, a, b):
    arrow(slide, *a, *b)


def textbox(slide, x, y, w, h, lines, size=10, color=TEXT):
    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tb.text_frame.word_wrap = True
    write(tb.text_frame, lines, size=size, color=color)
    return tb

# ---------------------------------------------------------------- conteúdo

USE_CASES = [
    ('Add a Playlist (Test Account)', [
        'CCPlayTV has no user account and no login of its own. The only credential is the IPTV playlist the user '
        'already has (Xtream Codes or M3U). It is typed on the TV and stays on the TV.',
        '1. On first launch the app opens "Connect your playlist". Otherwise, on "Select or Add your list", select '
        '"Add list" and press OK. The form opens on the "On the website" tab (see the portal use case).',
        '2. Move to "Xtream Codes" (or "M3U list") and fill in Playlist name, Host, User name and Password with the '
        'on-screen keyboard.',
        '3. Select "Connect and sync". The structure of the list is imported on the TV; when it finishes, select "Back".',
        '4. Select the new list card and press OK: the Home screen opens.',
        'Test playlist for the review (Xtream Codes):',
    ]),
    ('Trial and In-app Purchase', [
        'A new installation runs a free 7-day trial with every feature, counted on the TV — no account, voucher or '
        'payment needed. The reviewer tests the app inside this trial.',
        '1. Settings › "Portal & license" shows the license status ("Trial"), its end date and the offer: R$35 for '
        '12 months, for one TV, one-time purchase by Pix or card, no automatic renewal.',
        '2. "Buy license" opens "Buy a license with Pix": price, QR code and copy-and-paste code for the bank app. '
        'Paying by card is done on the website portal.ccplaytv.com (same TV ID and key).',
        '3. After paying, "Refresh" brings the license to the TV; the status changes to "Active" with the new end date.',
        '4. When the trial (or the license) ends, opening a list shows "A license is required to open the list", '
        'with "Buy license", "Refresh" and "Back". Lists, favorites and settings are kept; the app itself keeps '
        'opening.',
        'Payment is processed by Mercado Pago (Brazil). Nothing is charged automatically.',
    ]),
    ('Web Portal (Phone or PC)', [
        'Optional: the TV works completely without the portal.',
        '1. "Add list" › "On the website" (or Settings › "Portal & license") shows a QR code, the device ID and the '
        'device key.',
        '2. On the phone, scan the QR code (or open portal.ccplaytv.com and type the ID and key). The website shows '
        '"Lists", "License" and "Settings" for this TV.',
        '3. Add or edit a list on the website, then press "Reload"/"Refresh" on the TV: the list is imported on the TV.',
        '4. The TV talks to the portal only when the user presses one of these buttons; nothing runs in the '
        'background.',
    ]),
    ('Watch Live TV', [
        '1. On Home, move to "Live TV" in the top bar and press OK.',
        '2. The left column lists the categories of the playlist. Select one and press OK to see its channels.',
        '3. Each channel shows its number, logo and the current program when the playlist provides a guide (EPG). '
        'The right side shows the program now and the next ones.',
        '4. Press OK on a channel to watch it in full screen. CHANNEL UP/DOWN switches to the previous/next channel.',
        '5. RETURN closes the player and goes back to the channel, with the focus where it was.',
        'Moving the focus never starts playback: only OK does.',
    ]),
    ('Browse and Play Movies', [
        '1. On Home, select "Movies" in the top bar. Rows such as "Recently added" and "Top rated" are shown; '
        'LEFT opens the category list (★ Favorites, ↺ History, All and the categories of the playlist).',
        '2. Press OK on a category or on "See all" to open the grid. The grid has "Search" and "Sort".',
        '3. Press OK on a title to open its details: "Watch", "+ My List", "Similar", "Mark as watched", and the tabs '
        'Details / Cast / Similar.',
        '4. "Watch" starts the movie. In the player, OK shows the controls: -10s, Pause, +10s, Audio, Quality, '
        'Aspect and Info.',
        '5. A movie stopped in the middle shows "Resume from …" the next time.',
    ]),
    ('Watch a Series Episode', [
        '1. On Home, select "Series" in the top bar and open a category, as for movies.',
        '2. Press OK on a series. The details show the synopsis, "Watch T1:E1" (or "Continue" with progress), '
        'the season selector and the episode list.',
        '3. Choose a season, move to an episode and press OK to play it.',
        '4. RETURN goes back to the episode list, with the focus on the episode just watched.',
    ]),
    ('Favorites, My List and History', [
        '1. Hold OK, or press the YELLOW key, on a channel, movie, series or radio station to add it to Favorites '
        '(the same gesture removes it).',
        '2. Favorites appear in the "★ Favorites" category of each section and in the Home rows "My List" and '
        '"Favorite channels". Titles in progress appear in "Continue watching".',
        '3. Titles already played appear in "↺ History". In the History grid, the RED key removes the focused title; '
        'the details screen also has "Remove from history".',
    ]),
    ('Listen to Radio', [
        '1. Select "Radio" in the top bar. The left column has "★ Favorites", "Most played worldwide" and the countries.',
        '2. Select a country to list its stations (public directory Radio Browser; it does not use the playlist).',
        '3. Press OK on a station to play it; the "Now playing" bar shows it. STOP stops the radio.',
        '4. Hold OK or press YELLOW to keep the station in Favorites.',
    ]),
    ('Manage Lists, Categories and Parental Control', [
        '1. On "Select or Add your list", press DOWN on a list card: Resync, Edit, Delete and Configure.',
        '2. "Configure" opens the settings of that list: Account & list, EPG, Categories & parental control, History.',
        '3. In "Categories & parental control", categories can be hidden or locked with a PIN. Opening a locked '
        'category asks for the PIN (number keys or the on-screen keypad).',
        '4. "Settings" (bottom right of the list screen, or the gear in the top bar) has: Integrations (optional '
        'TMDB key of the user, for synopses and images), Playlists, Playback, Accessibility (including Language), '
        'About & credits, Portal & license and Backup (encrypted backup of favorites, history and progress).',
    ]),
]

# Cada tela: (arquivo, subtítulo, [(x_px, y_px, descrição)], nota opcional). (x, y) é o CENTRO do número na captura
# de 1920×1080 — num espaço livre ao lado do elemento, nunca em cima do texto que ele aponta.
SCREENS = [
    ('ui_01_lists.jpg', 'Select or Add your list', [
        (645, 560, 'List card (one per playlist). OK opens the list; DOWN shows Resync, Edit, Delete and Configure.'),
        (1255, 560, '"Add list": opens the playlist form.'),
        (1645, 1018, 'Settings for all lists.'),
        (230, 1018, '"Refresh": brings lists changed on the web portal.'),
    ], 'RETURN on this screen shows the "Exit CCPlayTV?" confirmation.'),
    ('ui_02_exit.jpg', 'Exit confirmation', [
        (1078, 560, '"Cancel": closes the dialog (default focus).'),
        (1203, 560, '"Exit": closes the application.'),
    ], None),
    ('ui_03a_add_site.jpg', 'Connect your playlist › On the website', [
        (1880, 409, 'Tabs: On the website (default), Xtream Codes, M3U list.'),
        (40, 690, 'QR code: opens the web portal for this TV on the phone.'),
        (990, 601, 'Portal address.'),
        (1880, 737, 'Device ID and device key, to sign in on the website.'),
        (380, 1050, '"Reload": imports lists added on the website. "Back": returns.'),
    ], None),
    ('ui_03_add_list.jpg', 'Connect your playlist › Xtream Codes', [
        (1880, 409, 'Type of list: Xtream Codes or M3U list (M3U asks for the list URL).'),
        (1880, 513, 'Playlist fields. OK on a field opens the on-screen keyboard.'),
        (1635, 803, '"Hide data": masks or shows the typed values.'),
        (1420, 882, '"Back": returns without saving.'),
        (1712, 960, '"Connect and sync": checks the list and imports it on the TV.'),
    ], None),
    ('ui_05_home.jpg', 'Home', [
        (1000, 71, 'Top bar: Home, Live TV, Movies, Series and Radio.'),
        (1585, 128, 'Active list: OK returns to the list selection.'),
        (1733, 128, 'Settings.'),
        (335, 143, 'Catalog status (last update of the list).'),
        (615, 734, 'Shortcuts to Live TV and Movies.'),
        (325, 817, 'Rows: Continue watching, My List, Favorite channels, Movies added, Series added.'),
        (400, 1026, 'Key guide of the current screen.'),
    ], None),
    ('ui_06_live.jpg', 'Live TV', [
        (220, 92, 'Current category and number of channels.'),
        (640, 93, 'Search this category and Full guide.'),
        (640, 315, 'Channel list: number, logo, name and program now (when the list has a guide).'),
        (700, 425, 'Focused channel: program now, time and progress.'),
        (668, 815, 'Next programs.'),
        (835, 1026, 'Key guide.'),
    ], 'OK on a channel plays it in full screen; hold OK or YELLOW adds it to Favorites.'),
    ('ui_07_live_player.jpg', 'Live TV player', [
        (575, 106, 'Channel banner: LIVE, number, logo, name and current program with progress.'),
    ], 'CHANNEL UP/DOWN switches channel. No seek bar on live channels: controls follow what the stream supports.'),
    ('ui_08_movies.jpg', 'Movies (Series has the same layout)', [
        (750, 71, 'Top bar with the current section.'),
        (96, 95, 'Category list: ★ Favorites, ↺ History, All and the categories of the playlist.'),
        (475, 152, 'Rows: Recently added, Top rated and others.'),
        (590, 612, 'Focused title with rating.'),
        (1735, 154, '"See all": opens the grid of the row.'),
    ], None),
    ('ui_09_category.jpg', 'Category grid', [
        (545, 74, 'Section, category and number of titles.'),
        (1397, 22, '"Search" in this category.'),
        (1722, 22, '"Sort": source order, A–Z, year, recently added, rating.'),
        (700, 700, 'Grid of titles. OK opens the details.'),
    ], None),
    ('ui_10_movie_detail.jpg', 'Movie details', [
        (472, 300, 'Title, year, duration and rating.'),
        (478, 419, 'Category and genres.'),
        (474, 598, '"Watch" (or "Resume from …").'),
        (786, 662, '"+ My List".'),
        (881, 662, 'Trailer (available with a TMDB key).'),
        (977, 662, '"Similar" titles.'),
        (1073, 662, '"Mark as watched".'),
        (445, 787, 'Tabs: Details, Cast, Similar.'),
    ], None),
    ('ui_11_movie_player.jpg', 'Movie / episode player', [
        (330, 94, 'Title.'),
        (560, 845, 'Position, progress and duration.'),
        (690, 1052, '-10s, Pause/Play, +10s.'),
        (1000, 1052, 'Audio and Quality (enabled when the stream has options).'),
        (1217, 1052, 'Aspect: Fit, Fill, Original, Zoom.'),
        (1346, 1052, 'Info about the title.'),
    ], None),
    ('ui_12_series_detail.jpg', 'Series details', [
        (955, 226, 'Title, year, seasons and rating.'),
        (657, 712, '"Watch T1:E1" (or "Continue" an episode in progress).'),
        (861, 570, '"+ My List".'),
        (1053, 570, '"Similar" titles.'),
        (135, 857, 'Season selector.'),
        (640, 914, 'Tabs: Similar, Cast, Details.'),
        (660, 1047, 'Episode list. OK plays the episode.'),
    ], None),
    ('ui_13_radio.jpg', 'Radio', [
        (330, 105, '"Now playing" bar.'),
        (300, 232, '★ Favorites, Most played worldwide and countries.'),
        (600, 208, 'Stations of the selected country. OK plays.'),
        (750, 1026, 'Key guide.'),
    ], 'Stations come from the public Radio Browser directory; STOP stops the radio.'),
    ('ui_04_settings.jpg', 'Settings', [
        (464, 216, 'Tabs: Integrations, Playlists, Playback, Accessibility, About & credits, Portal & license, Backup.'),
        (725, 345, 'A playlist: type, last sync, guide (EPG). "Configure" opens its settings.'),
        (715, 469, '"Add list".'),
    ], 'Integrations: optional TMDB key of the user. Accessibility: Reduce motion and Language.'),
    ('ui_04b_playback.jpg', 'Settings › Playback', [
        (900, 364, 'Default aspect.'),
        (900, 457, 'Default quality.'),
        (900, 550, 'Audio language.'),
        (900, 643, 'Subtitles.'),
    ], None),
    ('ui_04c_portal.jpg', 'Settings › Portal & license', [
        (885, 386, 'Access the portal: QR code, portal address, device ID and device key.'),
        (464, 563, 'Portal & license and Backup tabs.'),
        (1800, 467, 'License status ("Trial" or "Active"), end date and the offer: R$35 for 12 months.'),
        (1760, 636, '"Buy license" (Pix on the TV) and "Refresh" (brings the license after paying).'),
        (725, 860, '"Advanced": recovery code, new key and Unlink.'),
    ], None),
    ('ui_14_pix.jpg', 'Buy a license with Pix', [
        (40, 535, 'Pix QR code for the bank app.'),
        (1370, 317, 'Price and terms: R$35.00, 12 months, one TV, no automatic renewal.'),
        (1880, 603, 'Pix copy-and-paste code.'),
        (1340, 683, 'Card payment: on the website.'),
        (147, 900, '"Refresh": checks the payment and activates the license.'),
        (390, 831, '"Back": returns without paying.'),
    ], 'Payment processed by Mercado Pago. Screenshot with a fictional Pix code.'),
    ('ui_15_license_gate.jpg', 'License required (after the trial)', [
        (780, 245, 'Shown when a list is opened after the 7-day trial (or the license) ends.'),
        (40, 615, 'QR code, device ID and key: pay by card on the website.'),
        (169, 900, '"Buy license": Pix on the TV.'),
        (350, 900, '"Refresh": brings a license paid on the website.'),
        (496, 900, '"Back": returns to the list selection; lists and favorites are kept.'),
    ], None),
]

KEYS = [
    ('RETURN', 'Move to the previous screen; the focus goes back to the item you came from. On the list selection '
     'screen, shows the "Exit CCPlayTV?" confirmation.', 'Mandatory. Not modified.'),
    ('EXIT', 'Close the application and show the TV broadcast channel.', 'Mandatory. Not modified.'),
    ('UP / DOWN / LEFT / RIGHT', 'Move the focus between top bar, categories, rows, cards and buttons. Moving the '
     'focus never starts playback.', ''),
    ('ENTER (OK)', 'Open the focused item or press the focused button. On a channel or episode, start playback. '
     'In the player, show the controls.', ''),
    ('Hold ENTER (OK)', 'Add/remove the focused channel, movie, series or station to/from Favorites.', ''),
    ('YELLOW', 'Same as holding OK: add/remove Favorites.', 'Where items can be favorited'),
    ('RED', 'Remove the focused title from "↺ History".', 'History grid only'),
    ('PLAY / PAUSE', 'Play or pause the movie or episode.', 'Player'),
    ('STOP', 'Close the player. In Radio, stop the station.', 'Player, Radio'),
    ('REW / FF', 'Jump 10 seconds back / forward.', 'Only when the stream can seek'),
    ('CHANNEL UP / DOWN', 'Live TV player: previous / next channel. Live TV guide: previous / next page.', 'Live TV'),
    ('0 – 9', 'Type the parental-control PIN.', 'PIN prompt'),
]

LANGUAGE = [
    'The application is available in English, Portuguese and Spanish.',
    'By default it follows the TV language when it is one of these three (otherwise Portuguese).',
    'To switch the application language:',
    ('From "Select or Add your list", select "Settings" (or the gear in the top bar).', 1),
    ('Select "Accessibility", then "Language".', 1),
    ('Choose English, Português or Español. The new language applies immediately, no restart needed.', 1),
    'Names coming from the playlist (channels, categories, titles) are shown as the provider declares them and are '
    'not translated.',
]

# ---------------------------------------------------------------- montagem

def build(version):
    prs = Presentation(str(TEMPLATE))
    s = list(prs.slides)
    (instructions, cover, revision, contents1, ui_intro, ui_flow, ui_tree, ui_depth, contents2, uc_intro, uc_login,
     uc_menu, uc_iap, contents3, menus_intro, menus_numbers, menus_text, contents4, keys, contents5, language) = s

    # Capa e histórico
    write(placeholder(cover, 10).text_frame, ['CCPlayTV'], color=RGBColor(0xFF, 0xFF, 0xFF))
    write(placeholder(cover, 11).text_frame, ['Application UI Description', f'Version {version} · Samsung Smart TV (Tizen)'],
          size=22, color=RGBColor(0xFF, 0xFF, 0xFF))
    fill_table(shape_named(revision, 'Table 7'), ['Version', 'Date', 'Description', 'Author', 'Related App Version'],
               [['1.0', '2026.10.17', 'First UI description (first submission, commercial version)', 'John Costa', version]],
               size=10, header_size=11, col_widths=[0.8, 1.1, 3.5, 1.6, 2.0])

    # UI Structure: o slide "Flow Graph" recebe o diagrama real
    remove_shape(shape_named(ui_flow, 'Picture 6'))
    write(shape_named(ui_flow, '제목 1', placeholder_ok=False).text_frame, ['Flow Graph'], color=NAVY)
    write(placeholder(ui_flow, 1).text_frame, [
        'All screens and navigation paths. RETURN always goes back one level; the list chip in the top bar goes back '
        'to the list selection.'], size=12)
    placeholder(ui_flow, 1).height = Inches(0.7)
    draw_flow(ui_flow)

    # Use cases
    write(placeholder(uc_intro, 1).text_frame, [
        'CCPlayTV organizes and plays the IPTV playlist (Xtream Codes or M3U) that the user already has. It does not '
        'provide or sell any content, channel or playlist.',
        'No user account and no advertising. In-app purchase: a yearly license (R$35, 12 months, one TV) after a free '
        '7-day trial, paid by Pix on the TV or by card on the website — one-time, no automatic renewal.',
        'Use cases:',
        *[(f'{i}. {title}', 1) for i, (title, _) in enumerate(USE_CASES, start=1)],
    ], size=12)

    uc_slides = []
    title, steps = USE_CASES[0]
    placeholder(uc_login, 0).text_frame.text = 'Use Cases – Add a Playlist'
    remove_shape(shape_named(uc_login, 'Picture 1'))
    body = placeholder(uc_login, 1)
    body.top, body.height = Inches(1.6), Inches(3.35)
    write(body.text_frame, steps, size=12)
    login_table = shape_named(uc_login, 'Table 3')
    login_table.left, login_table.top = Inches(1.0), Inches(5.05)
    fill_table(login_table, ['Field', 'Value'], [
        ['Type', 'Xtream Codes'], ['Host', '<TEST_HOST>'], ['User name', '<TEST_USERNAME>'], ['Password', '<TEST_PASSWORD>'],
    ], size=10, header_size=10, col_widths=[1.6, 4.2])
    uc_slides.append(uc_login)
    uc_layout = uc_menu.slide_layout
    for title, steps in USE_CASES[1:]:
        slide = prs.slides.add_slide(uc_layout)
        clone_into(shape_named(uc_menu, 'Slide Number Placeholder 2'), slide)
        placeholder(slide, 0).text_frame.text = f'Use Cases – {title}'
        body = placeholder(slide, 1)
        body.left, body.top, body.width, body.height = Inches(0.5), Inches(1.75), Inches(9.0), Inches(5.1)
        write(body.text_frame, steps, size=13)
        uc_slides.append(slide)

    # Menus and Functions
    write(placeholder(menus_intro, 1).text_frame, [
        'One slide per screen, with numbered labels on a real screenshot of this version.',
        'Screenshots taken with the application in English and a fictional demo playlist (names, posters and logos '
        'invented for the demo). Content names come from the playlist and are not translated.',
        'Every screen keeps one item focused (orange frame); the key guide at the bottom shows the keys of that screen.',
    ], size=14)
    screen_slides = [screen_slide(prs, menus_numbers, *screen) for screen in SCREENS]

    # Remote keys: duas páginas da tabela do template
    keys_page2 = prs.slides.add_slide(keys.slide_layout)
    clone_into(shape_named(keys, 'Slide Number Placeholder 2'), keys_page2)
    placeholder(keys_page2, 0).text_frame.text = 'Remote Control Keys'
    remove_shape(placeholder(keys_page2, 1))
    write(placeholder(keys, 1).text_frame, [
        'Only the keys below are used. RETURN and EXIT keep their standard behavior.'], size=13)
    placeholder(keys, 1).height = Inches(0.8)
    table = shape_named(keys, 'Table 8')
    table.top = Inches(2.35)
    clone = clone_into(table, keys_page2)
    clone.top = Inches(1.75)
    widths = [1.85, 4.85, 1.95]
    fill_table(table, ['Key', 'Action', 'Remarks'], [list(k) for k in KEYS[:6]], size=9.5, header_size=10, col_widths=widths)
    fill_table(clone, ['Key', 'Action', 'Remarks'], [list(k) for k in KEYS[6:]], size=9.5, header_size=10, col_widths=widths)

    write(placeholder(language, 1).text_frame, LANGUAGE, size=14)

    # Fora: instruções, exemplos e diagramas de exemplo
    for slide in (instructions, ui_intro, ui_tree, ui_depth, uc_menu, uc_iap, menus_numbers, menus_text):
        delete_slide(prs, slide)

    reorder(prs, [cover, revision, contents1, ui_flow, contents2, uc_intro, *uc_slides, contents3, menus_intro,
                  *screen_slides, contents4, keys, keys_page2, contents5, language])
    out = KIT / f'CCPlayTV_UI_Description_{version}.pptx'
    prs.save(str(out))
    return out


def screen_slide(prs, example, image, subtitle, labels, note):
    slide = prs.slides.add_slide(example.slide_layout)
    clone_into(shape_named(example, 'Slide Number Placeholder 2'), slide)
    placeholder(slide, 0).text_frame.text = 'Menus and Functions'
    remove_shape(placeholder(slide, 1))
    sub = clone_into(shape_named(example, '제목 1', placeholder_ok=False), slide)
    write(sub.text_frame, [subtitle], color=NAVY)
    pic_left, pic_top, pic_w = 0.35, 1.95, 6.1
    pic_h = pic_w * 9 / 16
    picture_fit(slide, KIT / image, pic_left, pic_top, pic_w)
    oval_src = shape_named(example, 'Oval 8')
    d = Emu(oval_src.width).inches
    for n, (x, y, _) in enumerate(labels, start=1):
        oval = clone_into(oval_src, slide)
        cx = pic_left + x / 1920 * pic_w
        cy = pic_top + y / 1080 * pic_h
        oval.left = Inches(min(max(cx - d / 2, pic_left), pic_left + pic_w - d))
        oval.top = Inches(min(max(cy - d / 2, pic_top), pic_top + pic_h - d))
        oval.text_frame.paragraphs[0].runs[0].text = str(n)
    table = clone_into(shape_named(example, 'Table 3'), slide)
    table.left, table.top = Inches(6.6), Inches(1.95)
    fill_table(table, ['No.', 'Description'], [[str(n), text] for n, (_, _, text) in enumerate(labels, start=1)],
               size=8.5, header_size=9, col_widths=[0.45, 2.6])
    if note:
        textbox(slide, pic_left, pic_top + pic_h + 0.12, pic_w, 0.8, [note], size=10)
    return slide


def draw_flow(slide):
    y0 = 2.35
    splash = box(slide, 0.25, y0 + 1.75, 1.15, 0.55, 'Splash', 'app launch')
    add = box(slide, 1.7, y0, 1.6, 0.6, 'Connect your playlist', 'On the website (QR) · Xtream Codes · M3U')
    progress = box(slide, 1.7, y0 + 0.85, 1.6, 0.5, 'Import progress')
    lists = box(slide, 1.7, y0 + 1.7, 1.6, 0.65, 'Select or Add your list', 'one card per playlist')
    exit_ = box(slide, 1.7, y0 + 2.75, 1.6, 0.5, 'Exit confirmation', 'RETURN on the list screen')
    settings = box(slide, 3.65, y0, 1.85, 0.75, 'Settings', 'Integrations · Playlists · Playback · Accessibility · About · Portal & license · Backup')
    actions = box(slide, 3.65, y0 + 0.95, 1.85, 0.6, 'List actions', 'Resync · Edit · Delete · Configure')
    home = box(slide, 3.65, y0 + 1.75, 1.85, 0.6, 'Home', 'top bar + rows')
    live = box(slide, 5.85, y0, 1.75, 0.6, 'Live TV', 'categories → channels + guide')
    movies = box(slide, 5.85, y0 + 0.8, 1.75, 0.6, 'Movies', 'rows → category grid')
    series = box(slide, 5.85, y0 + 1.6, 1.75, 0.6, 'Series', 'rows → category grid')
    radio = box(slide, 5.85, y0 + 2.4, 1.75, 0.6, 'Radio', 'countries → stations')
    gate = box(slide, 3.65, y0 + 2.6, 1.85, 0.6, 'License required', 'list opened after the 7-day trial')
    box(slide, 3.65, y0 + 3.45, 1.85, 0.5, 'PIN prompt', 'locked category (Live TV, Movies, Series)')
    pix = box(slide, 5.85, y0 + 3.3, 1.75, 0.6, 'Pix purchase', 'R$35 · 12 months · one TV')
    live_player = box(slide, 7.95, y0, 1.75, 0.6, 'Live player', 'CH +/- switches channel')
    movie_detail = box(slide, 7.95, y0 + 0.8, 1.75, 0.6, 'Movie details', 'Watch → full-screen player')
    series_detail = box(slide, 7.95, y0 + 1.6, 1.75, 0.6, 'Series details', 'season → episode → player')
    now_playing = box(slide, 7.95, y0 + 2.4, 1.75, 0.6, 'Now playing bar', 'station playing')

    link(slide, right(splash), left(lists))
    link(slide, right(splash), left(add))
    link(slide, bottom(add), top(progress))
    link(slide, bottom(progress), top(lists))
    link(slide, bottom(lists), top(exit_))
    link(slide, right(lists), left(settings))
    link(slide, right(lists), left(actions))
    link(slide, right(lists), left(home))
    for section in (live, movies, series, radio):
        link(slide, right(home), left(section))
    link(slide, right(live), left(live_player))
    link(slide, right(movies), left(movie_detail))
    link(slide, right(series), left(series_detail))
    link(slide, right(radio), left(now_playing))
    link(slide, right(lists), left(gate))
    link(slide, right(gate), left(pix))
    textbox(slide, 0.25, 6.3, 9.4, 0.6, [
        'Settings opens from the list screen or the gear in the top bar; Settings › Portal & license also opens the Pix '
        'purchase. List actions appear with DOWN on a list card; a locked category asks for the PIN. The website '
        'portal.ccplaytv.com (QR code on the TV) manages lists, license and settings from a phone or PC.'], size=9)


if __name__ == '__main__':
    parser = argparse.ArgumentParser()
    parser.add_argument('--version', default='0.2.1')
    print(build(parser.parse_args().version))
